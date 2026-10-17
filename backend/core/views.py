import base64
import hmac
import io
import logging
import os
import random
import threading
from datetime import datetime, timedelta

import cloudinary.uploader
import openpyxl
from openpyxl.styles import Font, PatternFill
import qrcode
import sendgrid
import stripe
from django.conf import settings
from django.contrib.auth.password_validation import validate_password
from django.contrib.auth.tokens import PasswordResetTokenGenerator
from django.core import signing
from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import IntegrityError, connection, transaction
from django.urls import reverse
from django.utils.encoding import force_bytes, force_str
from django.utils.http import urlsafe_base64_decode, urlsafe_base64_encode
from sendgrid.helpers.mail import (
    Mail, Email, To, Content,
    Attachment, FileContent, FileName, FileType, Disposition,
)
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable, Flowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_RIGHT, TA_CENTER
from reportlab.lib.colors import HexColor
from rest_framework import viewsets, status
from rest_framework.views import APIView
from rest_framework.decorators import api_view, authentication_classes, permission_classes
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, AllowAny
from rest_framework.exceptions import PermissionDenied
from rest_framework_simplejwt.tokens import RefreshToken
from django.contrib.auth import authenticate, get_user_model
from django.db.models import Count, Max, Q, Sum
from django.http import HttpResponse
from django.template.response import TemplateResponse
from django.utils import timezone
from boutique.pricing import OrderPricingError, price_order, to_cents
from .models import (
    BARBER_CHOICES, DEFAULT_BARBER,
    Barbershop, Barber, Service, Reservation, User, ClientMedia, ClosedPeriod, OpeningHours, Review, WaitingList,
    BarberLeave, Notification, NotificationPreference, PaymentCard, UserSession,
)
from .emails import (
    format_date_fr, format_date_fr_short, format_heure_fr, is_late_cancellation, queue_reservation_email,
)
from .oauth import InvalidOAuthToken, OAuthNotConfigured, verify_google_id_token
from .permissions import IsStaffRole
from .recommendations import SALON_TZ, build_recommendation, create_client_notification, run_daily_recommendations
from .slots import day_hours, minutes_to_hhmm, slot_times_for, weekly_hours
from .sms import send_reminders_for_tomorrow
from .serializers import (
    BarbershopSerializer, BarberSerializer, ServiceSerializer, ReservationSerializer,
    UserSerializer, RegisterSerializer, ClientMediaSerializer,
    ClosedPeriodSerializer, OpeningHoursSerializer, ReviewSerializer, WaitingListSerializer,
    BarberLeaveSerializer, NotificationSerializer,
    NotificationPreferenceSerializer, PaymentCardSerializer,
)

logger = logging.getLogger(__name__)

ACOMPTE_FIXE = 5
CANCELLED_STATUSES = ('cancelled', 'cancelled_client')


class PayeStamp(Flowable):
    def __init__(self, label='PAYÉ'):
        super().__init__()
        self.label = label

    def draw(self):
        self.canv.saveState()
        self.canv.translate(60, 30)
        self.canv.rotate(-15)
        stamp_color = HexColor('#7BAF9E')
        self.canv.setStrokeColor(stamp_color)
        self.canv.setFillColor(stamp_color)
        self.canv.setLineWidth(2)
        self.canv.roundRect(-45, -18, 90, 36, 4, stroke=1, fill=0)
        self.canv.setFont('Helvetica-Bold', 22)
        self.canv.drawCentredString(0, -8, self.label)
        self.canv.restoreState()

    def wrap(self, *args):
        return (120, 60)


class SoldeRegleStamp(Flowable):
    def draw(self):
        self.canv.saveState()
        self.canv.translate(60, 30)
        self.canv.rotate(-15)
        stamp_color = HexColor('#2D6A4F')
        self.canv.setStrokeColor(stamp_color)
        self.canv.setFillColor(stamp_color)
        self.canv.setLineWidth(2)
        self.canv.roundRect(-62, -18, 124, 36, 4, stroke=1, fill=0)
        self.canv.setFont('Helvetica-Bold', 15)
        self.canv.drawCentredString(0, -6, 'SOLDE RÉGLÉ')
        self.canv.restoreState()

    def wrap(self, *args):
        return (140, 60)


def fmt_eur(n):
    return f"{n:.2f}".replace('.', ',') + ' €'


class BarbershopViewSet(viewsets.ModelViewSet):
    queryset = Barbershop.objects.all()
    serializer_class = BarbershopSerializer


class BarberViewSet(viewsets.ModelViewSet):
    queryset = Barber.objects.all()
    serializer_class = BarberSerializer

    def get_permissions(self):
        # La liste de l'équipe est publique (affichée côté client dans certains
        # écrans) mais la création/modification/suppression reste réservée à
        # Willo, comme pour ServiceViewSet ci-dessous.
        if self.action in ('list', 'retrieve'):
            return [AllowAny()]
        return [IsStaffRole()]


class ServiceViewSet(viewsets.ModelViewSet):
    queryset = Service.objects.all()
    serializer_class = ServiceSerializer

    def get_permissions(self):
        # La liste des prestations doit être publique (un client non connecté doit
        # pouvoir consulter les prix avant de réserver) — mais la création/modification/
        # suppression reste réservée à Willo, comme pour ClosedPeriodViewSet plus bas.
        if self.action in ('list', 'retrieve'):
            return [AllowAny()]
        return [IsStaffRole()]


class ClosedPeriodViewSet(viewsets.ModelViewSet):
    queryset = ClosedPeriod.objects.all()
    serializer_class = ClosedPeriodSerializer
    http_method_names = ['get', 'post', 'delete', 'head', 'options']

    def get_permissions(self):
        if self.action in ('list', 'retrieve'):
            return [AllowAny()]
        return [IsStaffRole()]

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)


class BarberLeaveViewSet(viewsets.ModelViewSet):
    queryset = BarberLeave.objects.all()
    serializer_class = BarberLeaveSerializer
    http_method_names = ['get', 'post', 'delete', 'head', 'options']

    def get_permissions(self):
        if self.action in ('list', 'retrieve'):
            return [AllowAny()]
        return [IsStaffRole()]


class NotificationViewSet(viewsets.ModelViewSet):
    # Notifications du coiffeur uniquement : celles des clients ont un user.
    queryset = Notification.objects.filter(user__isnull=True)
    serializer_class = NotificationSerializer
    permission_classes = [IsStaffRole]
    http_method_names = ['get', 'patch', 'head', 'options']


class ClientNotificationViewSet(viewsets.ModelViewSet):
    serializer_class = NotificationSerializer
    permission_classes = [IsAuthenticated]
    http_method_names = ['get', 'patch', 'head', 'options']

    def get_queryset(self):
        return Notification.objects.filter(user=self.request.user)


class WaitingListViewSet(viewsets.ModelViewSet):
    queryset = WaitingList.objects.all()
    serializer_class = WaitingListSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        is_staff = user.is_superuser or getattr(user, 'role', None) == 'admin'
        if is_staff:
            return WaitingList.objects.all()
        return WaitingList.objects.filter(client=user)

    def perform_create(self, serializer):
        serializer.save(client=self.request.user)


@api_view(['GET', 'PUT'])
@permission_classes([AllowAny])
def opening_hours(request):
    """Horaires hebdomadaires du salon (7 jours, lundi = 0).

    GET : public (le flow de réservation en a besoin). PUT : staff uniquement,
    liste de jours [{weekday, is_closed, open_time, close_time}] — seuls les
    jours envoyés sont modifiés."""
    if request.method == 'PUT':
        if not IsStaffRole().has_permission(request, None):
            raise PermissionDenied('Réservé au salon.')
        entries = request.data if isinstance(request.data, list) else [request.data]
        errors = {}
        validated = []
        for entry in entries:
            weekday = entry.get('weekday') if isinstance(entry, dict) else None
            if weekday not in range(7):
                return Response({'error': 'weekday (0-6) requis pour chaque jour.'}, status=status.HTTP_400_BAD_REQUEST)
            instance = OpeningHours.objects.filter(weekday=weekday).first()
            serializer = OpeningHoursSerializer(instance, data=entry, partial=instance is not None)
            if serializer.is_valid():
                validated.append(serializer)
            else:
                errors[weekday] = serializer.errors
        if errors:
            return Response(errors, status=status.HTTP_400_BAD_REQUEST)
        with transaction.atomic():
            for serializer in validated:
                serializer.save()

    hours = OpeningHours.objects.all()
    return Response(OpeningHoursSerializer(hours, many=True).data)


@api_view(['GET'])
@permission_classes([AllowAny])
def available_slots(request):
    date_str = request.query_params.get('date')
    if not date_str:
        return Response({'error': 'date requis (YYYY-MM-DD).'}, status=status.HTTP_400_BAD_REQUEST)
    try:
        requested_date = datetime.strptime(date_str, '%Y-%m-%d').date()
    except ValueError:
        return Response({'error': 'Format de date invalide (YYYY-MM-DD attendu).'}, status=status.HTTP_400_BAD_REQUEST)

    # Agenda du barbier choisi (prénom, cf. BARBER_CHOICES) ; sans paramètre,
    # celui du barbier par défaut (anciennes versions de l'app).
    barber = request.query_params.get('barber') or DEFAULT_BARBER
    if barber not in dict(BARBER_CHOICES):
        return Response({'error': 'Barbier inconnu.'}, status=status.HTTP_400_BAD_REQUEST)

    on_leave = BarberLeave.objects.filter(
        barber=barber, start_date__lte=requested_date, end_date__gte=requested_date,
    ).exists()
    # Jour fermé (horaires hebdo du salon ou fermeture exceptionnelle) : aucun créneau.
    span = day_hours(requested_date)
    is_closed = span is None or ClosedPeriod.objects.filter(
        start_date__lte=requested_date, end_date__gte=requested_date,
    ).exists()
    booked_hhmm = {
        t.strftime('%H:%M')
        for t in Reservation.objects
            .filter(barber=barber, date=requested_date)
            .exclude(status__in=['cancelled', 'cancelled_client'])
            .values_list('time', flat=True)
    }

    slots = [
        {'start_time': slot, 'end_time': slot, 'is_available': not on_leave and slot not in booked_hhmm}
        for slot in ([] if is_closed else slot_times_for(requested_date))
    ]

    return Response({
        'date': date_str,
        'barber': barber,
        'is_closed': is_closed,
        'open_time': minutes_to_hhmm(span[0]) if span else None,
        'close_time': minutes_to_hhmm(span[1] % (24 * 60)) if span else None,
        'slots': slots,
    })


@api_view(['GET'])
@permission_classes([IsStaffRole])
def barbers_week_availability(request):
    """Créneaux encore libres par barbier et par jour, sur une semaine
    (?week_start=YYYY-MM-DD, lundi de la semaine en cours par défaut).

    État d'un jour : 'closed' (jour fermé dans les horaires ou fermeture
    exceptionnelle du salon), 'leave' (congé du
    barbier), 'past' (jour écoulé) ou 'open'. Aujourd'hui, seuls les créneaux
    pas encore commencés comptent comme libres."""
    now = timezone.localtime(timezone.now(), SALON_TZ)
    today = now.date()
    try:
        week_start = datetime.strptime(request.query_params['week_start'], '%Y-%m-%d').date()
    except (KeyError, ValueError):
        week_start = today - timedelta(days=today.weekday())
    days = [week_start + timedelta(days=i) for i in range(7)]

    closed_periods = list(ClosedPeriod.objects.filter(start_date__lte=days[-1], end_date__gte=days[0]))
    leaves = list(BarberLeave.objects.filter(start_date__lte=days[-1], end_date__gte=days[0]))
    active = Reservation.objects.exclude(status__in=CANCELLED_STATUSES)
    booked = {}
    for barber, d, t in active.filter(date__range=(days[0], days[-1])).values_list('barber', 'date', 'time'):
        booked.setdefault((barber, d), set()).add(t.strftime('%H:%M'))
    totals = dict(active.values_list('barber').annotate(n=Count('id')).values_list('barber', 'n'))
    hours = weekly_hours()
    grids = {d: slot_times_for(d, hours) for d in days}

    def day_state(barber, d):
        taken = booked.get((barber, d), set())
        info = {'date': d.isoformat(), 'total': len(grids[d]), 'booked': len(taken), 'free': 0}
        period = next((p for p in closed_periods if p.start_date <= d <= p.end_date), None)
        if hours[d.weekday()] is None or period:
            return {**info, 'state': 'closed', 'reason': period.reason if period else 'Fermé'}
        if any(l.barber == barber and l.start_date <= d <= l.end_date for l in leaves):
            return {**info, 'state': 'leave', 'reason': 'Congé'}
        if d < today:
            return {**info, 'state': 'past'}
        upcoming = [s for s in grids[d] if d > today or s > now.strftime('%H:%M')]
        return {**info, 'state': 'open', 'free': sum(s not in taken for s in upcoming)}

    return Response({
        'week_start': week_start.isoformat(),
        'barbers': {
            barber: {
                'total_rdv': totals.get(barber, 0),
                'week_rdv': sum(len(booked.get((barber, d), ())) for d in days),
                'days': [day_state(barber, d) for d in days],
            }
            for barber, _ in BARBER_CHOICES
        },
    })


password_reset_token_generator = PasswordResetTokenGenerator()


def send_password_reset_email(request, user):
    try:
        sg = sendgrid.SendGridAPIClient(api_key=os.environ.get('SENDGRID_API_KEY'))
        uidb64 = urlsafe_base64_encode(force_bytes(user.pk))
        token = password_reset_token_generator.make_token(user)
        reset_url = request.build_absolute_uri(
            reverse('password-reset-confirm-page', args=[uidb64, token])
        )
        html_content = f"""
        <div style="font-family: Georgia, serif; background-color: #0D0C0A; color: #FFFFFF; padding: 40px; max-width: 600px; margin: auto;">
            <h1 style="color: #C9A84C; font-size: 32px; margin-bottom: 4px;">WilloBarber</h1>
            <p style="color: #6B6560; font-size: 12px; letter-spacing: 2px; margin-top: 0;">SALON DE BARBIER · BRUXELLES</p>
            <hr style="border: 1px solid #1A1814; margin: 24px 0;">
            <h2 style="color: #FFFFFF;">Réinitialisation du mot de passe</h2>
            <p style="color: #CCCCCC;">Vous avez demandé à réinitialiser votre mot de passe. Cliquez sur le bouton ci-dessous pour en choisir un nouveau.</p>
            <div style="text-align: center; margin: 32px 0;">
                <a href="{reset_url}" style="background: #C9A84C; color: #1A1208; text-decoration: none; font-weight: bold; padding: 14px 28px; border-radius: 100px; display: inline-block;">Réinitialiser mon mot de passe</a>
            </div>
            <p style="color: #6B6560; font-size: 12px;">Ce lien est valable 3 jours. Si vous n'êtes pas à l'origine de cette demande, ignorez simplement cet email.</p>
            <hr style="border: 1px solid #1A1814; margin: 24px 0;">
            <p style="color: #C9A84C; font-size: 14px; text-align: center;">À bientôt chez WilloBarber ✂️</p>
        </div>
        """
        message = Mail(
            from_email=Email(settings.SENDGRID_FROM_EMAIL, settings.SENDGRID_FROM_NAME),
            to_emails=To(user.email),
            subject='🔑 Réinitialisation de votre mot de passe WilloBarber',
            html_content=Content('text/html', html_content)
        )
        response = sg.send(message)
        print(f"[EMAIL] Reset mot de passe envoyé à {user.email} — SendGrid status={response.status_code}")
        return True
    except Exception as e:
        print(f"[EMAIL ERROR] Échec envoi reset mot de passe à {user.email}: {e}")
        return False


def _reset_user_password(uidb64, token, new_password):
    """Vérifie l'uid/token et applique le nouveau mot de passe.

    Retourne (user, None) en cas de succès, (None, message_erreur) sinon.
    """
    try:
        uid = force_str(urlsafe_base64_decode(uidb64))
        user = User.objects.get(pk=uid)
    except (TypeError, ValueError, OverflowError, User.DoesNotExist):
        return None, "Lien de réinitialisation invalide."

    if not password_reset_token_generator.check_token(user, token):
        return None, "Ce lien de réinitialisation est invalide ou a expiré."

    try:
        validate_password(new_password, user=user)
    except DjangoValidationError as e:
        return None, ' '.join(e.messages)

    user.set_password(new_password)
    user.save()
    return user, None


class PasswordResetRequestView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        email = (request.data.get('email') or '').strip()
        generic_response = Response(
            {'detail': "Si un compte existe avec cet email, un lien de réinitialisation vient d'être envoyé."}
        )
        if not email:
            return Response({'error': 'email requis'}, status=400)

        user = User.objects.filter(email__iexact=email, is_active=True).first()
        if user:
            send_password_reset_email(request, user)
        return generic_response


class PasswordResetConfirmView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        uidb64 = request.data.get('uid')
        token = request.data.get('token')
        new_password = request.data.get('new_password')
        new_password2 = request.data.get('new_password2')

        if not uidb64 or not token or not new_password:
            return Response({'error': 'uid, token et new_password sont requis.'}, status=400)
        if new_password2 is not None and new_password != new_password2:
            return Response({'error': 'Les mots de passe ne correspondent pas.'}, status=400)

        user, error = _reset_user_password(uidb64, token, new_password)
        if error:
            return Response({'error': error}, status=400)
        return Response({'detail': 'Mot de passe réinitialisé avec succès.'})


def _password_reset_page_html(*, uidb64, token, error=None, success=False):
    if success:
        body = """
            <h2 style="color:#FFFFFF;">Mot de passe mis à jour ✂️</h2>
            <p style="color:#CCCCCC;">Votre mot de passe a bien été réinitialisé. Vous pouvez maintenant vous reconnecter dans l'app WilloBarber.</p>
        """
    else:
        error_html = f'<div style="background:rgba(192,57,43,0.15);border:1px solid #C0392B;border-radius:10px;padding:12px;margin-bottom:20px;color:#FF6B6B;font-size:14px;">{error}</div>' if error else ''
        body = f"""
            <h2 style="color:#FFFFFF;margin-bottom:8px;">Nouveau mot de passe</h2>
            <p style="color:#CCCCCC;margin-bottom:20px;">Choisissez votre nouveau mot de passe ci-dessous.</p>
            {error_html}
            <form method="post">
                <label style="display:block;font-size:11px;letter-spacing:1px;color:#C9A84C;text-transform:uppercase;margin-bottom:6px;">Nouveau mot de passe</label>
                <input type="password" name="new_password" required minlength="8"
                    style="width:100%;box-sizing:border-box;background:rgba(255,255,255,0.05);border:1.5px solid rgba(255,255,255,0.12);border-radius:12px;padding:14px 16px;color:#FFFFFF;font-size:15px;margin-bottom:16px;" />
                <label style="display:block;font-size:11px;letter-spacing:1px;color:#C9A84C;text-transform:uppercase;margin-bottom:6px;">Confirmer le mot de passe</label>
                <input type="password" name="new_password2" required minlength="8"
                    style="width:100%;box-sizing:border-box;background:rgba(255,255,255,0.05);border:1.5px solid rgba(255,255,255,0.12);border-radius:12px;padding:14px 16px;color:#FFFFFF;font-size:15px;margin-bottom:20px;" />
                <button type="submit"
                    style="width:100%;background:#C9A84C;color:#1A1208;font-weight:bold;font-size:15px;border:none;border-radius:100px;padding:15px;cursor:pointer;">
                    Réinitialiser mon mot de passe
                </button>
            </form>
        """
    return f"""<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="utf-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1" />
    <title>WilloBarber — Réinitialisation du mot de passe</title>
</head>
<body style="margin:0;background:#0D0C0A;font-family:Georgia,serif;">
    <div style="max-width:440px;margin:60px auto;padding:32px;background:#1A1814;border-radius:24px;">
        <h1 style="color:#C9A84C;font-size:26px;margin:0 0 2px 0;">WilloBarber</h1>
        <p style="color:#6B6560;font-size:11px;letter-spacing:2px;margin:0 0 20px 0;">SALON DE BARBIER · BRUXELLES</p>
        <hr style="border:1px solid #252018;margin:0 0 24px 0;" />
        {body}
    </div>
</body>
</html>"""


def password_reset_confirm_page(request, uidb64, token):
    if request.method == 'POST':
        new_password = request.POST.get('new_password') or ''
        new_password2 = request.POST.get('new_password2') or ''
        if new_password != new_password2:
            html = _password_reset_page_html(uidb64=uidb64, token=token, error="Les mots de passe ne correspondent pas.")
            return HttpResponse(html)

        user, error = _reset_user_password(uidb64, token, new_password)
        if error:
            html = _password_reset_page_html(uidb64=uidb64, token=token, error=error)
            return HttpResponse(html)

        return HttpResponse(_password_reset_page_html(uidb64=uidb64, token=token, success=True))

    return HttpResponse(_password_reset_page_html(uidb64=uidb64, token=token))


class ReservationViewSet(viewsets.ModelViewSet):
    queryset = Reservation.objects.all()
    serializer_class = ReservationSerializer

    def get_queryset(self):
        user = self.request.user
        is_staff = user.is_superuser or getattr(user, 'role', None) == 'admin'
        if is_staff:
            return Reservation.objects.all()
        return Reservation.objects.filter(user=user)

    # Champs qu'un client n'est jamais autorisé à modifier lui-même — seul le
    # staff (via IsStaffRole) peut les toucher, sinon n'importe quel client
    # pourrait falsifier son propre statut de paiement.
    STAFF_ONLY_FIELDS = {'payment_status', 'payment_method', 'amount_paid'}
    # Seules ces actions restent ouvertes à un client sur sa propre réservation :
    # l'annuler (status='cancelled_client' + motif) ou la reprogrammer (date/heure)
    # — cf. reservationsApi.cancel()/.reschedule() côté mobile.
    CLIENT_ALLOWED_FIELDS = {'status', 'cancellation_reason', 'date', 'time'}

    def _app_url(self):
        return settings.APP_URL or self.request.build_absolute_uri('/')

    def perform_update(self, serializer):
        user = self.request.user
        is_staff = user.is_superuser or getattr(user, 'role', None) == 'admin'
        if not is_staff:
            requested_fields = set(self.request.data.keys())
            if requested_fields & self.STAFF_ONLY_FIELDS:
                raise PermissionDenied("Seul le salon peut modifier ces informations.")
            if requested_fields - self.CLIENT_ALLOWED_FIELDS:
                raise PermissionDenied("Vous ne pouvez qu'annuler ou reprogrammer votre réservation.")
            # Un client ne peut faire évoluer le statut que vers "annulé par le client".
            if 'status' in requested_fields and self.request.data.get('status') != 'cancelled_client':
                raise PermissionDenied("Vous ne pouvez qu'annuler votre réservation.")
        previous_status = serializer.instance.status
        previous_slot = {'date': serializer.instance.date, 'time': serializer.instance.time}
        reservation = serializer.save()

        # Emails client : annulation (client ou salon) ou reprogrammation.
        if previous_status not in CANCELLED_STATUSES and reservation.status in CANCELLED_STATUSES:
            queue_reservation_email(
                reservation, 'cancelled', self._app_url(), late=is_late_cancellation(reservation),
            )
        elif reservation.status not in CANCELLED_STATUSES and (
            (reservation.date, reservation.time) != (previous_slot['date'], previous_slot['time'])
        ):
            queue_reservation_email(reservation, 'rescheduled', self._app_url(), previous=previous_slot)

        # Retirer les points gagnés à la création pour éviter le farming
        # (annulation par le client ou par le salon, une seule fois).
        if previous_status in ('pending', 'confirmed') and reservation.status in ('cancelled', 'cancelled_client'):
            points_to_refund = int(round(reservation.service.price * 2))
            client = reservation.user
            client.loyalty_points = max(0, (client.loyalty_points or 0) - points_to_refund)
            client.save(update_fields=['loyalty_points'])

    def create(self, request, *args, **kwargs):
        date = request.data.get('date')
        time = request.data.get('time')
        conflict_response = Response(
            {'error': "Ce créneau vient d'être pris. Veuillez choisir un autre horaire."},
            status=status.HTTP_409_CONFLICT,
        )

        # Un client ne réserve que dans la grille du jour (horaires du salon) ;
        # le staff garde la main pour caler un RDV hors horaires.
        user = request.user
        if not (user.is_superuser or getattr(user, 'role', None) == 'admin'):
            try:
                requested_date = datetime.strptime(str(date), '%Y-%m-%d').date()
            except ValueError:
                requested_date = None
            if requested_date and (
                ClosedPeriod.objects.filter(start_date__lte=requested_date, end_date__gte=requested_date).exists()
                or str(time)[:5] not in slot_times_for(requested_date)
            ):
                return Response(
                    {'error': "Le salon est fermé à cet horaire. Veuillez choisir un autre créneau."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

        with transaction.atomic():
            # select_for_update ferme la fenêtre de course quand une réservation
            # active existe déjà sur ce créneau. Elle ne verrouille rien sur un
            # créneau tout neuf (rien à sélectionner) — c'est la contrainte
            # unique_active_reservation_slot en base (cf. models.py) qui ferme
            # ce second cas : deux requêtes simultanées sur un créneau libre
            # provoquent un IntegrityError pour l'une des deux, rattrapé plus bas.
            conflict = Reservation.objects.select_for_update().filter(
                barber=request.data.get('barber') or DEFAULT_BARBER,
                date=date,
                time=time,
                status__in=['pending', 'confirmed'],
            ).exists()

            if conflict:
                return conflict_response

            try:
                # Savepoint imbriqué : si super().create() lève une IntegrityError,
                # seul ce savepoint est annulé — la transaction englobante reste
                # utilisable pour renvoyer une réponse propre au lieu de planter.
                with transaction.atomic():
                    return super().create(request, *args, **kwargs)
            except IntegrityError:
                return conflict_response
            except Exception as e:
                print(f"[RESERVATION] Bad Request — data reçue: {request.data} — erreur: {e}")
                raise

    def perform_create(self, serializer):
        print(f"[RESERVATION] data: {self.request.data}")
        print(f"[RESERVATION] user: {self.request.user}")
        reservation = serializer.save(user=self.request.user)
        print(f"[RESERVATION] sauvegardée: id={reservation.id}")

        # Paiement Stripe confirmé côté app avant la création (cf. book.tsx) :
        # seule la 1re réservation d'un groupe enfants porte le payment_intent_id.
        if reservation.payment_intent_id:
            site_root = self.request.build_absolute_uri('/')
            transaction.on_commit(lambda: threading.Thread(
                target=send_online_payment_invoice,
                args=(reservation.id, site_root),
                daemon=True,
            ).start())

        # Attribution automatique des points fidélité : 2 pts par euro du service
        points_earned = int(round(reservation.service.price * 2))
        user = self.request.user
        user.loyalty_points = (user.loyalty_points or 0) + points_earned
        user.save(update_fields=['loyalty_points'])

        queue_reservation_email(reservation, 'confirmed', self._app_url())


class RegisterView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        serializer = RegisterSerializer(data=request.data)
        if serializer.is_valid():
            user = serializer.save()
            refresh = RefreshToken.for_user(user)
            return Response({
                'user': UserSerializer(user).data,
                'refresh': str(refresh),
                'access': str(refresh.access_token),
            }, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class LoginView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        username = request.data.get('username')
        password = request.data.get('password')
        user = authenticate(request=request, username=username, password=password)
        if user:
            refresh = RefreshToken.for_user(user)
            access = str(refresh.access_token)
            refresh_str = str(refresh)
            if user.role == 'admin':
                UserSession.objects.create(
                    user=user,
                    device=request.META.get('HTTP_USER_AGENT', '')[:255],
                    access_token=access,
                    refresh_token=refresh_str,
                )
            return Response({
                'user': UserSerializer(user).data,
                'refresh': refresh_str,
                'access': access,
            })
        return Response(
            {'detail': 'Identifiants invalides.', 'debug': {'username_received': username, 'password_empty': not password}},
            status=status.HTTP_401_UNAUTHORIZED,
        )


class LogoutView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        refresh_str = request.data.get('refresh')
        try:
            token = RefreshToken(refresh_str)
            token.blacklist()
        except Exception:
            return Response(status=status.HTTP_400_BAD_REQUEST)
        UserSession.objects.filter(user=request.user, refresh_token=refresh_str).delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class MeView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response(UserSerializer(request.user).data)

    def patch(self, request):
        serializer = UserSerializer(request.user, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class ChangePasswordView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        user = request.user
        if not user.check_password(request.data.get('old_password')):
            return Response({'detail': 'Ancien mot de passe incorrect.'}, status=status.HTTP_400_BAD_REQUEST)
        user.set_password(request.data.get('new_password'))
        user.save()
        return Response({'detail': 'Mot de passe mis à jour.'})


class DeleteAccountView(APIView):
    """Suppression RGPD initiée par le client (droit à l'effacement).

    Soft delete anonymisé plutôt que suppression réelle de la ligne User :
    Reservation.user est en CASCADE, une vraie suppression détruirait
    l'historique de réservations passées qu'on doit garder pour la
    comptabilité. Même mécanisme que l'archivage staff (delete_client) mais
    avec anonymisation des données personnelles et is_active=False pour
    bloquer toute reconnexion.
    """
    permission_classes = [IsAuthenticated]

    def delete(self, request):
        user = request.user

        Reservation.objects.filter(
            user=user,
            date__gte=timezone.localdate(),
            status__in=['pending', 'confirmed'],
        ).update(status='cancelled_client')

        timestamp = timezone.now().strftime('%Y%m%d%H%M%S%f')
        user.first_name = 'Client supprimé'
        user.last_name = ''
        user.email = f'deleted_{timestamp}_{user.pk}@willobarber.fr'
        user.phone = ''
        user.is_deleted = True
        user.deleted_at = timezone.now()
        user.deletion_reason = 'client'
        user.is_active = False
        user.save(update_fields=[
            'first_name', 'last_name', 'email', 'phone',
            'is_deleted', 'deleted_at', 'deletion_reason', 'is_active',
        ])

        UserSession.objects.filter(user=user).delete()
        PaymentCard.objects.filter(owner=user).delete()

        refresh_str = request.data.get('refresh')
        if refresh_str:
            try:
                RefreshToken(refresh_str).blacklist()
            except Exception:
                pass

        return Response(status=status.HTTP_200_OK)


class OAuthLoginView(APIView):
    """Connexion Google : l'app envoie l'id_token signé par Google, vérifié
    ici (cf. core/oauth.py). L'email et le nom sont lus dans le token vérifié —
    jamais dans le body, sinon n'importe qui pourrait se connecter au compte
    d'un autre en connaissant son email. Crée le compte client à la volée si
    l'email n'existe pas encore (sans mot de passe utilisable)."""
    permission_classes = [AllowAny]

    def post(self, request):
        if request.data.get('provider', 'google') != 'google':
            return Response({'error': 'Fournisseur non pris en charge.'}, status=400)
        id_token = (request.data.get('id_token') or '').strip()
        if not id_token:
            return Response({'error': 'id_token requis.'}, status=400)
        try:
            payload = verify_google_id_token(id_token)
        except OAuthNotConfigured:
            return Response({'error': 'Connexion Google non configurée.'}, status=503)
        except InvalidOAuthToken as e:
            logger.warning('[OAUTH] id_token Google refusé : %s', e)
            return Response({'error': 'Connexion Google invalide.'}, status=401)

        email = payload['email'].strip().lower()
        first_name = (payload.get('given_name') or '').strip()
        last_name = (payload.get('family_name') or '').strip()

        user = User.objects.filter(email__iexact=email).first()
        if not user:
            base_username = email.split('@')[0] or 'client'
            username = base_username
            suffix = 1
            while User.objects.filter(username=username).exists():
                suffix += 1
                username = f'{base_username}{suffix}'
            user = User.objects.create_user(
                username=username,
                email=email,
                first_name=first_name,
                last_name=last_name,
            )
            user.set_unusable_password()
            user.save()

        refresh = RefreshToken.for_user(user)
        return Response({
            'access': str(refresh.access_token),
            'refresh': str(refresh),
            'user': {
                'username': user.username,
                'email': user.email,
                'first_name': user.first_name,
            },
        })


class CheckClientView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        identifier = request.data.get('identifier', '').strip()
        print(f"DEBUG CHECK-CLIENT: identifier='{identifier}'")

        all_users = User.objects.all()
        print(f"DEBUG USERS COUNT: {all_users.count()}")
        for u in all_users[:3]:
            print(f"DEBUG USER: username={u.username} email={u.email}")

        if not identifier:
            return Response({'error': 'identifier requis'}, status=400)

        if '@' in identifier:
            user = User.objects.filter(email__iexact=identifier).first()
        else:
            digits = ''.join(filter(str.isdigit, identifier))
            user = User.objects.filter(phone=digits).first() if digits else None

        print(f"DEBUG FOUND: {user}")
        if user:
            return Response({'status': 'exists', 'first_name': user.first_name or user.username})
        return Response({'status': 'new'})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def recommendations(request):
    # ?exclude=2026-10-24T14:00,2026-10-25T15:30 : créneaux déjà proposés au
    # client (bouton « Voir d'autres créneaux » de RecommendationModal).
    exclude = [key for key in request.query_params.get('exclude', '').split(',') if key]
    return Response(build_recommendation(request.user, exclude_slots=exclude))


@api_view(['POST'])
@authentication_classes([])
@permission_classes([AllowAny])
def run_recommendations_task(request):
    """Appelé une fois par jour par cron-job.org avec le header X-CRON-TOKEN."""
    expected = settings.CRON_TOKEN
    if not expected:
        return Response({'error': 'CRON_TOKEN non configuré.'}, status=status.HTTP_503_SERVICE_UNAVAILABLE)
    if not hmac.compare_digest(request.headers.get('X-CRON-TOKEN', ''), expected):
        return Response({'error': 'Token invalide.'}, status=status.HTTP_403_FORBIDDEN)
    return Response(run_daily_recommendations())


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def reservation_qr_code(request, pk):
    try:
        reservation = Reservation.objects.get(pk=pk)
    except Reservation.DoesNotExist:
        return Response({'error': 'Réservation introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    is_staff = request.user.is_superuser or getattr(request.user, 'role', None) == 'admin'
    if not is_staff and reservation.user_id != request.user.id:
        return Response({'error': "Vous n'avez pas accès à ce QR code."}, status=status.HTTP_403_FORBIDDEN)

    return Response({
        'reservation_id': reservation.id,
        'qr_code': str(reservation.qr_code),
        'checked_in': reservation.checked_in,
    })


@api_view(['POST'])
@permission_classes([IsStaffRole])
def checkin_reservation(request, qr_code):
    try:
        reservation = Reservation.objects.select_related('user', 'service').get(qr_code=qr_code)
    except Reservation.DoesNotExist:
        return Response({'error': 'QR code invalide — aucune réservation correspondante.'}, status=status.HTTP_404_NOT_FOUND)

    already_checked_in = reservation.checked_in
    if not already_checked_in:
        reservation.checked_in = True
        reservation.checked_in_at = timezone.now()
        reservation.save(update_fields=['checked_in', 'checked_in_at'])

    client = reservation.user
    client_name = f"{client.first_name} {client.last_name}".strip() or client.username

    return Response({
        'success': True,
        'already_checked_in': already_checked_in,
        'reservation_id': reservation.id,
        'client_name': client_name,
        'client_email': client.email,
        'client_phone': client.phone,
        'service_name': reservation.service.name,
        'date': reservation.date,
        'time': reservation.time,
        'checked_in_at': reservation.checked_in_at,
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def generate_invoice(request, pk):
    try:
        reservation = Reservation.objects.get(pk=pk)
    except Reservation.DoesNotExist:
        return Response({'error': 'Réservation introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    is_staff = request.user.is_superuser or getattr(request.user, 'role', None) == 'admin'
    if not is_staff and reservation.user_id != request.user.id:
        return Response({'error': "Vous n'avez pas accès à cette facture."}, status=status.HTTP_403_FORBIDDEN)

    print(f"[FACTURE] Génération pour réservation {reservation.id} - {reservation.service.name} - {reservation.service.price}€")
    invoice_number = f"WB-2026-{random.randint(10000, 99999)}"
    pdf_bytes = build_online_invoice_pdf(reservation, invoice_number)

    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="Facture_WilloBarber_{invoice_number}.pdf"'
    return response


def build_online_invoice_pdf(reservation, invoice_number):
    """Facture du montant réglé en ligne (acompte ou totalité)."""
    user = reservation.user
    service = reservation.service
    today = format_date_fr_short(datetime.now().date())

    quantity = getattr(reservation, 'child_count', 1) or 1
    total_ttc = float(service.price) * quantity
    acompte = float(reservation.amount_paid) if reservation.amount_paid is not None else ACOMPTE_FIXE
    is_fully_paid = acompte >= total_ttc
    sous_total_ht = total_ttc / 1.21
    tva = total_ttc - sous_total_ht
    solde = 0 if is_fully_paid else max(total_ttc - acompte, 0)
    methode_label = reservation.get_payment_method_display() or 'Carte'

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
        rightMargin=20*mm, leftMargin=20*mm,
        topMargin=20*mm, bottomMargin=20*mm)

    gold = colors.HexColor('#C9A84C')
    dark = colors.HexColor('#1A1814')
    cream = colors.HexColor('#F5F0E8')
    green = colors.HexColor('#2D6A4F')
    grey = colors.HexColor('#888888')

    styles = getSampleStyleSheet()

    elements = []

    # Header
    header_data = [[
        Paragraph('<b>w willobarber</b><br/><font size=8 color=grey>SALON DE BARBIER · BERCHEM</font>', styles['Normal']),
        Paragraph(f'<font size=26><i>Facture.</i></font><br/><br/><font size=10 color=grey>N° {invoice_number}</font><br/><font size=10 color=grey>Émise le {today}</font>', ParagraphStyle('h', alignment=TA_RIGHT, spaceAfter=6, leading=32))
    ]]
    header_table = Table(header_data, colWidths=[90*mm, 90*mm])
    header_table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP')]))
    elements.append(header_table)
    elements.append(HRFlowable(width="100%", thickness=1, color=dark))
    elements.append(Spacer(1, 10*mm))

    # Parties
    first_name = user.first_name or user.username
    last_name = user.last_name or ''
    parties_data = [[
        Paragraph(f'<font size=9 color=grey>ÉMETTEUR</font><br/><b>WilloBarber SRL</b><br/>78 rue Auguste van Zande<br/>1082 Bruxelles · Belgique<br/>contact@willobarber.be<br/>+32 2 555 04 04<br/><font size=9 color=grey>TVA BE 0789.123.456</font>', styles['Normal']),
        [
            Paragraph(f'<font size=9 color=grey>FACTURÉ À</font><br/><b>{first_name} {last_name}</b><br/>{user.email}', ParagraphStyle('right', alignment=TA_RIGHT)),
            PayeStamp(label='RÉGLÉ' if is_fully_paid else 'PAYÉ'),
        ]
    ]]
    parties_table = Table(parties_data, colWidths=[90*mm, 90*mm])
    parties_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('ALIGN', (1,0), (1,0), 'RIGHT'),
    ]))
    elements.append(parties_table)
    elements.append(Spacer(1, 10*mm))

    # Tableau prestations
    unit_price_ttc = float(service.price)
    unit_price_ht = unit_price_ttc / 1.21
    data = [
        ['DESCRIPTION', 'QTÉ', 'P.U. HT', 'TVA', 'TOTAL TTC'],
        [service.name, str(quantity), fmt_eur(unit_price_ht), '21%', fmt_eur(total_ttc)],
    ]
    t = Table(data, colWidths=[80*mm, 20*mm, 30*mm, 20*mm, 30*mm])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), dark),
        ('TEXTCOLOR', (0,0), (-1,0), gold),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 10),
        ('ROWBACKGROUND', (0,1), (-1,-1), cream),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#dddddd')),
        ('PADDING', (0,0), (-1,-1), 8),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 8*mm))

    # Récap
    recap_data = [
        ['Méthode', methode_label],
        ['Sous-total HT', fmt_eur(sous_total_ht)],
        ['TVA 21%', fmt_eur(tva)],
        ['Total TTC', fmt_eur(total_ttc)],
        ['INTÉGRALEMENT RÉGLÉ' if is_fully_paid else 'Acompte réglé', f'-{fmt_eur(acompte)}'],
        ['Solde dû au salon', fmt_eur(solde)],
    ]
    recap_table = Table(recap_data, colWidths=[80*mm, 40*mm], hAlign='RIGHT')
    recap_table.setStyle(TableStyle([
        ('FONTSIZE', (0,0), (-1,-1), 10),
        ('TEXTCOLOR', (1,3), (1,3), dark),
        ('FONTNAME', (0,3), (-1,3), 'Helvetica-Bold'),
        ('FONTNAME', (0,5), (-1,5), 'Helvetica-Bold'),
        ('TEXTCOLOR', (0,4), (-1,4), green),
        ('LINEABOVE', (0,3), (-1,3), 1, dark),
        ('LINEABOVE', (0,5), (-1,5), 1, dark),
        ('ALIGN', (1,0), (1,-1), 'RIGHT'),
        ('PADDING', (0,0), (-1,-1), 6),
    ]))
    elements.append(recap_table)
    elements.append(Spacer(1, 6*mm))

    # Bloc RÉGLÉ CE JOUR
    regle_label = 'RÉGLÉ CE JOUR · Paiement intégral' if is_fully_paid else 'RÉGLÉ CE JOUR · Acompte de réservation'
    regle_data = [[regle_label, fmt_eur(acompte)]]
    regle_table = Table(regle_data, colWidths=[120*mm, 40*mm], hAlign='RIGHT')
    regle_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), dark),
        ('TEXTCOLOR', (0,0), (0,0), colors.white),
        ('TEXTCOLOR', (1,0), (1,0), gold),
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 11),
        ('ALIGN', (1,0), (1,0), 'RIGHT'),
        ('PADDING', (0,0), (-1,-1), 10),
        ('ROUNDEDCORNERS', [6, 6, 6, 6]),
    ]))
    elements.append(regle_table)
    elements.append(Spacer(1, 6*mm))
    elements.append(HRFlowable(width="100%", thickness=0.5, color=grey))
    elements.append(Spacer(1, 4*mm))
    note_text = (
        f"Cette prestation a été réglée intégralement en ligne ({fmt_eur(acompte)}). Aucun solde ne reste dû au salon."
        if is_fully_paid else
        f"Cette facture concerne uniquement l'acompte de {fmt_eur(acompte)} encaissé pour sécuriser la réservation. Une facture finale sera émise à l'issue du rendez-vous."
    )
    elements.append(Paragraph(
        note_text,
        ParagraphStyle('note', fontSize=9, textColor=grey)
    ))

    doc.build(elements)
    buffer.seek(0)
    return buffer.read()


@api_view(['GET'])
@permission_classes([IsStaffRole])
def export_reservations_pdf(request):
    reservations = Reservation.objects.select_related('user', 'service').order_by('-date', '-time')

    gold = colors.HexColor('#C9A84C')
    dark = colors.HexColor('#1A1814')
    cream = colors.HexColor('#F5F0E8')
    grey = colors.HexColor('#888888')

    now = timezone.localtime(timezone.now())

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
        rightMargin=15*mm, leftMargin=15*mm,
        topMargin=20*mm, bottomMargin=15*mm)

    elements = [
        Paragraph('<font size=20><b>Rapport WilloBarber</b></font>', getSampleStyleSheet()['Normal']),
        Paragraph(
            f'<font size=9 color=grey>Généré le {format_date_fr(now.date())} à {now.strftime("%H:%M")}</font>',
            getSampleStyleSheet()['Normal'],
        ),
        Spacer(1, 8*mm),
    ]

    data = [['Date', 'Heure', 'Client', 'Service', 'Statut', 'Montant']]
    for r in reservations:
        client_name = r.user.get_full_name() or r.user.username
        data.append([
            r.date.strftime('%d/%m/%Y'),
            r.time.strftime('%H:%M'),
            client_name,
            r.service.name,
            r.get_status_display(),
            fmt_eur(float(r.service.price)),
        ])

    table = Table(data, colWidths=[24*mm, 16*mm, 40*mm, 40*mm, 30*mm, 30*mm], repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), dark),
        ('TEXTCOLOR', (0, 0), (-1, 0), gold),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8.5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, cream]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#dddddd')),
        ('PADDING', (0, 0), (-1, -1), 6),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer.read(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="rapport_willobarber.pdf"'
    return response


@api_view(['GET'])
@permission_classes([IsStaffRole])
def export_reservations_excel(request):
    reservations = Reservation.objects.select_related('user', 'service').order_by('-date', '-time')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Réservations'

    headers = ['Date', 'Heure', 'Client', 'Service', 'Statut', 'Montant']
    ws.append(headers)
    header_fill = PatternFill(start_color='1A1814', end_color='1A1814', fill_type='solid')
    header_font = Font(color='C9A84C', bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for r in reservations:
        client_name = r.user.get_full_name() or r.user.username
        ws.append([
            r.date.strftime('%d/%m/%Y'),
            r.time.strftime('%H:%M'),
            client_name,
            r.service.name,
            r.get_status_display(),
            float(r.service.price),
        ])

    for col in ws.columns:
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col) + 2
        ws.column_dimensions[col[0].column_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="willobarber.xlsx"'
    return response


def get_or_create_stripe_customer(user):
    """Customer Stripe du client, créé (et mémorisé sur User) au 1er besoin."""
    stripe.api_key = settings.STRIPE_SECRET_KEY
    if user.stripe_customer_id:
        try:
            customer = stripe.Customer.retrieve(user.stripe_customer_id)
            if not getattr(customer, 'deleted', False):
                return customer
        except stripe.error.InvalidRequestError:
            # Customer supprimé côté Stripe ou créé avec une autre clé (test ↔ live).
            pass
    customer = stripe.Customer.create(
        email=user.email or None,
        name=f"{user.first_name} {user.last_name}".strip() or user.username,
        metadata={'user_id': user.id},
    )
    user.stripe_customer_id = customer.id
    user.save(update_fields=['stripe_customer_id'])
    return customer


def _customer_payment_method(user, payment_method_id):
    """La carte, si elle appartient bien au customer de ce client — sinon None."""
    if not user.stripe_customer_id:
        return None
    try:
        payment_method = stripe.PaymentMethod.retrieve(payment_method_id)
    except stripe.error.InvalidRequestError:
        return None
    return payment_method if payment_method.customer == user.stripe_customer_id else None


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def create_payment_intent(request):
    """Deux cas :

    - Nouvelle carte (pas de payment_method) : PaymentIntent rattaché au
      customer avec setup_future_usage='on_session' — l'app le confirme avec la
      carte saisie, et Stripe la mémorise sur le customer une fois réglé.
    - Carte enregistrée (payment_method=pm_…) : PaymentIntent confirmé ici,
      directement (off_session=True, confirm=True) → status 'succeeded', rien à
      saisir côté app.

    3D Secure : même avec une carte enregistrée, la banque peut exiger une
    nouvelle authentification (SCA). Stripe lève alors une CardError
    'authentication_required' ; on renvoie le client_secret du PaymentIntent
    avec status 'requires_action' et l'app le confirme avec la même carte, ce
    qui affiche le challenge 3DS (cf. StripeCardField.*.tsx).

    clientSecret reste toujours présent dans la réponse : les versions de l'app
    antérieures aux cartes enregistrées ne lisent que ce champ.
    """
    stripe.api_key = settings.STRIPE_SECRET_KEY
    amount = request.data.get('amount')
    currency = request.data.get('currency', 'eur')
    payment_method_id = request.data.get('payment_method')

    metadata = {'user_id': request.user.id, 'username': request.user.username}

    # Commande boutique : le montant est recalculé ici depuis les prix en base,
    # exactement comme à la création de la commande — jamais celui du client.
    order = request.data.get('order')
    if order is not None:
        if not isinstance(order, dict):
            return Response({'error': 'Commande invalide.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            _, _, total = price_order(
                order.get('items') or [],
                order.get('delivery_type') or 'retrait_salon',
                (order.get('delivery_address') or '').strip(),
            )
        except OrderPricingError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        print(f"[PAYMENT] create boutique user={request.user.id}: client_amount={amount}, "
              f"total_calculated={total}, stripe_amount={to_cents(total)}")
        if currency != 'eur':
            return Response({'error': 'Devise invalide.'}, status=status.HTTP_400_BAD_REQUEST)
        amount = to_cents(total)
        metadata.update(kind='boutique', total=str(total))
    else:
        try:
            amount = int(round(float(amount) * 100))
        except (TypeError, ValueError):
            return Response({'error': 'Montant invalide.'}, status=status.HTTP_400_BAD_REQUEST)
        print(f"[PAYMENT] create user={request.user.id}: stripe_amount={amount}")

    if amount <= 0:
        return Response({'error': 'Montant invalide.'}, status=status.HTTP_400_BAD_REQUEST)

    params = {
        'amount': amount,
        'currency': currency,
        'metadata': metadata,
    }

    if payment_method_id:
        if not _customer_payment_method(request.user, payment_method_id):
            return Response({'error': 'Carte introuvable.'}, status=status.HTTP_404_NOT_FOUND)
        try:
            intent = stripe.PaymentIntent.create(
                **params,
                customer=request.user.stripe_customer_id,
                payment_method=payment_method_id,
                off_session=True,
                confirm=True,
            )
        except stripe.error.CardError as e:
            pending = getattr(e.error, 'payment_intent', None)
            if e.code == 'authentication_required' and pending:
                return Response({
                    'clientSecret': pending['client_secret'],
                    'status': 'requires_action',
                    'paymentIntentId': pending['id'],
                })
            return Response({'error': e.user_message or str(e)}, status=status.HTTP_400_BAD_REQUEST)
        except stripe.error.StripeError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        return Response({'clientSecret': intent.client_secret, 'status': intent.status, 'paymentIntentId': intent.id})

    try:
        customer = get_or_create_stripe_customer(request.user)
        params.update(customer=customer.id, setup_future_usage='on_session')
    except stripe.error.StripeError as e:
        # La mémorisation de carte est un bonus : le paiement passe sans customer.
        print(f"[STRIPE] Customer indisponible pour user {request.user.id}: {e}")

    try:
        intent = stripe.PaymentIntent.create(**params)
    except stripe.error.StripeError as e:
        return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

    return Response({'clientSecret': intent.client_secret, 'status': intent.status, 'paymentIntentId': intent.id})


def _saved_cards(user):
    """Cartes du customer, une par empreinte : payer deux fois avec la même carte
    « nouvelle » l'enregistre deux fois chez Stripe (la plus récente gagne)."""
    if not user.stripe_customer_id or not settings.STRIPE_SECRET_KEY:
        return {}
    stripe.api_key = settings.STRIPE_SECRET_KEY
    by_fingerprint = {}
    for pm in stripe.PaymentMethod.list(customer=user.stripe_customer_id, type='card').auto_paging_iter():
        by_fingerprint.setdefault(pm.card.fingerprint, []).append(pm)
    return by_fingerprint


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def my_payment_methods(request):
    try:
        cards = _saved_cards(request.user)
    except stripe.error.StripeError as e:
        print(f"[STRIPE] Liste des cartes indisponible pour user {request.user.id}: {e}")
        cards = {}
    return Response([
        {
            'id': pms[0].id,
            'brand': pms[0].card.brand,
            'last4': pms[0].card.last4,
            'exp_month': pms[0].card.exp_month,
            'exp_year': pms[0].card.exp_year,
        }
        for pms in cards.values()
    ])


@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def delete_my_payment_method(request, pm_id):
    try:
        cards = _saved_cards(request.user)
        duplicates = next((pms for pms in cards.values() if any(pm.id == pm_id for pm in pms)), None)
        if duplicates is None:
            return Response({'error': 'Carte introuvable.'}, status=status.HTTP_404_NOT_FOUND)
        for pm in duplicates:
            stripe.PaymentMethod.detach(pm.id)
    except stripe.error.StripeError as e:
        return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)
    return Response(status=status.HTTP_204_NO_CONTENT)


def build_final_invoice_pdf(reservation):
    """Facture finale émise quand le solde est réglé au salon (cash/carte)."""
    user = reservation.user
    service = reservation.service
    invoice_number = f"WB-FINAL-{reservation.id}"
    today = format_date_fr_short(datetime.now().date())

    quantity = getattr(reservation, 'child_count', 1) or 1
    total_ttc = float(service.price) * quantity
    acompte = ACOMPTE_FIXE
    solde = float(reservation.amount_paid) if reservation.amount_paid is not None else max(total_ttc - acompte, 0)
    methode_label = reservation.get_payment_method_display() or '—'

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
        rightMargin=20*mm, leftMargin=20*mm,
        topMargin=20*mm, bottomMargin=20*mm)

    gold = colors.HexColor('#C9A84C')
    dark = colors.HexColor('#1A1814')
    green = colors.HexColor('#2D6A4F')
    grey = colors.HexColor('#888888')

    styles = getSampleStyleSheet()
    elements = []

    # Header
    header_data = [[
        Paragraph('<b>w willobarber</b><br/><font size=8 color=grey>SALON DE BARBIER · BERCHEM</font>', styles['Normal']),
        Paragraph(f'<font size=26><i>Facture finale.</i></font><br/><br/><font size=10 color=grey>N° {invoice_number}</font><br/><font size=10 color=grey>Émise le {today}</font>', ParagraphStyle('h', alignment=TA_RIGHT, spaceAfter=6, leading=32))
    ]]
    header_table = Table(header_data, colWidths=[90*mm, 90*mm])
    header_table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP')]))
    elements.append(header_table)
    elements.append(HRFlowable(width="100%", thickness=1, color=dark))
    elements.append(Spacer(1, 10*mm))

    # Parties
    first_name = user.first_name or user.username
    last_name = user.last_name or ''
    parties_data = [[
        Paragraph(f'<font size=9 color=grey>ÉMETTEUR</font><br/><b>WilloBarber SRL</b><br/>78 rue Auguste van Zande<br/>1082 Bruxelles · Belgique<br/>contact@willobarber.be<br/>+32 2 555 04 04<br/><font size=9 color=grey>TVA BE 0789.123.456</font>', styles['Normal']),
        [
            Paragraph(f'<font size=9 color=grey>FACTURÉ À</font><br/><b>{first_name} {last_name}</b><br/>{user.email}', ParagraphStyle('right', alignment=TA_RIGHT)),
            SoldeRegleStamp(),
        ]
    ]]
    parties_table = Table(parties_data, colWidths=[90*mm, 90*mm])
    parties_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('ALIGN', (1,0), (1,0), 'RIGHT'),
    ]))
    elements.append(parties_table)
    elements.append(Spacer(1, 10*mm))

    # Récapitulatif
    recap_data = [
        [service.name, fmt_eur(total_ttc)],
        ['Acompte déjà réglé en ligne', f'-{fmt_eur(acompte)}'],
        [f'Solde réglé au salon ({methode_label})', fmt_eur(solde)],
        ['Total réglé', fmt_eur(total_ttc)],
    ]
    recap_table = Table(recap_data, colWidths=[130*mm, 40*mm], hAlign='RIGHT')
    recap_table.setStyle(TableStyle([
        ('FONTSIZE', (0,0), (-1,-1), 10.5),
        ('TEXTCOLOR', (1,1), (1,1), green),
        ('TEXTCOLOR', (1,2), (1,2), green),
        ('FONTNAME', (0,3), (-1,3), 'Helvetica-Bold'),
        ('FONTSIZE', (0,3), (-1,3), 12.5),
        ('TEXTCOLOR', (1,3), (1,3), gold),
        ('LINEABOVE', (0,3), (-1,3), 1, dark),
        ('ALIGN', (1,0), (1,-1), 'RIGHT'),
        ('PADDING', (0,0), (-1,-1), 8),
    ]))
    elements.append(recap_table)
    elements.append(Spacer(1, 8*mm))
    elements.append(HRFlowable(width="100%", thickness=0.5, color=grey))
    elements.append(Spacer(1, 4*mm))
    elements.append(Paragraph(
        "Prestation entièrement réglée. Merci de votre confiance.",
        ParagraphStyle('note', fontSize=10, textColor=green, alignment=TA_CENTER)
    ))

    doc.build(elements)
    buffer.seek(0)
    return buffer.read(), invoice_number


@api_view(['POST'])
@permission_classes([IsStaffRole])
def send_final_invoice(request, pk):
    try:
        reservation = Reservation.objects.get(pk=pk)
    except Reservation.DoesNotExist:
        return Response({'error': 'Réservation introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    if reservation.payment_status != 'paid_onsite':
        return Response(
            {'error': "La réservation n'est pas marquée comme payée au salon."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    pdf_bytes, invoice_number = build_final_invoice_pdf(reservation)

    try:
        send_invoice_email(
            reservation.user,
            subject=f'✂️ Facture finale — {reservation.service.name}',
            heading=f"Merci pour votre visite {reservation.user.first_name or reservation.user.username} ✂️",
            body="Votre prestation est entièrement réglée. Vous trouverez votre facture finale en pièce jointe.",
            pdf_bytes=pdf_bytes,
            filename=f'{invoice_number}.pdf',
        )
        print(f"[EMAIL] Facture finale envoyée à {reservation.user.email}")
    except Exception as e:
        print(f"[EMAIL ERROR] Facture finale: {str(e)}")
        return Response(
            {'error': "Le paiement a été enregistré mais l'email n'a pas pu être envoyé."},
            status=status.HTTP_502_BAD_GATEWAY,
        )

    return Response({'success': True, 'invoice_number': invoice_number})


def send_invoice_email(user, subject, heading, body, pdf_bytes, filename):
    sg = sendgrid.SendGridAPIClient(api_key=os.environ.get('SENDGRID_API_KEY'))
    message = Mail(
        from_email=Email(settings.SENDGRID_FROM_EMAIL, settings.SENDGRID_FROM_NAME),
        to_emails=To(user.email),
        subject=subject,
        html_content=Content('text/html', f"""
            <div style="font-family: Georgia, serif; background-color: #0D0C0A; color: #FFFFFF; padding: 40px; max-width: 600px; margin: auto;">
                <h1 style="color: #C9A84C; font-size: 32px; margin-bottom: 4px;">WilloBarber</h1>
                <p style="color: #6B6560; font-size: 12px; letter-spacing: 2px; margin-top: 0;">SALON DE BARBIER · BRUXELLES</p>
                <hr style="border: 1px solid #1A1814; margin: 24px 0;">
                <h2 style="color: #FFFFFF;">{heading}</h2>
                <p style="color: #CCCCCC;">{body}</p>
                <hr style="border: 1px solid #1A1814; margin: 24px 0;">
                <p style="color: #C9A84C; font-size: 14px; text-align: center;">À bientôt chez WilloBarber ✂️</p>
            </div>
        """),
    )
    message.attachment = Attachment(
        FileContent(base64.b64encode(pdf_bytes).decode()),
        FileName(filename),
        FileType('application/pdf'),
        Disposition('attachment'),
    )
    sg.send(message)


# Lien « Télécharger la facture » de la notification client : ouvert hors de
# l'app (navigateur), donc sans JWT — l'accès est porté par une signature.
INVOICE_LINK_SALT = 'online-invoice-download'
INVOICE_LINK_MAX_AGE = 60 * 60 * 24 * 365


def online_invoice_number(reservation):
    return f"WB-{reservation.date.year}-{reservation.id:05d}"


def _payment_intent_succeeded(payment_intent_id):
    if not settings.STRIPE_SECRET_KEY:
        print(f"[FACTURE] STRIPE_SECRET_KEY absent — PaymentIntent {payment_intent_id} non vérifié")
        return True
    stripe.api_key = settings.STRIPE_SECRET_KEY
    return stripe.PaymentIntent.retrieve(payment_intent_id).status == 'succeeded'


def send_online_payment_invoice(reservation_id, site_root):
    """Après un paiement en ligne : PDF → email au client → notification
    « invoice » (SSE). Tourne dans un thread après le commit de la réservation,
    pour ne pas rallonger la réponse de POST /reservations/."""
    try:
        reservation = Reservation.objects.select_related('user', 'service').get(pk=reservation_id)
        if not _payment_intent_succeeded(reservation.payment_intent_id):
            print(f"[FACTURE] PaymentIntent non réglé pour réservation {reservation_id} — pas de facture")
            return

        invoice_number = online_invoice_number(reservation)
        pdf_bytes = build_online_invoice_pdf(reservation, invoice_number)
        token = signing.dumps(reservation.id, salt=INVOICE_LINK_SALT)
        invoice_url = site_root.rstrip('/') + reverse('online-invoice-download', args=[token])
        amount = float(reservation.amount_paid or 0)

        try:
            send_invoice_email(
                reservation.user,
                subject=f'🧾 Votre facture WilloBarber — {reservation.service.name}',
                heading=f"Merci {reservation.user.first_name or reservation.user.username} ✂️",
                body=f"Votre paiement de {fmt_eur(amount)} est confirmé. Vous trouverez votre facture en pièce jointe.",
                pdf_bytes=pdf_bytes,
                filename=f'Facture_WilloBarber_{invoice_number}.pdf',
            )
            print(f"[FACTURE] {invoice_number} envoyée à {reservation.user.email}")
        except Exception as e:
            print(f"[EMAIL ERROR] Facture {invoice_number}: {str(e)}")

        create_client_notification(
            reservation.user, 'invoice',
            title='Votre facture est disponible',
            desc=f"{reservation.service.name} · {format_date_fr_short(reservation.date)} · {fmt_eur(amount)}",
            data={
                'reservation_id': reservation.id,
                'invoice_number': invoice_number,
                'invoice_url': invoice_url,
                'amount_paid': amount,
            },
        )
    except Exception as e:
        import traceback
        print(f"[FACTURE ERROR] réservation {reservation_id}: {str(e)}")
        print(traceback.format_exc())
    finally:
        connection.close()


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def download_online_invoice(request, token):
    try:
        reservation_id = signing.loads(token, salt=INVOICE_LINK_SALT, max_age=INVOICE_LINK_MAX_AGE)
        reservation = Reservation.objects.select_related('user', 'service').get(pk=reservation_id)
    except (signing.BadSignature, Reservation.DoesNotExist):
        return Response({'error': 'Lien de facture invalide ou expiré.'}, status=status.HTTP_404_NOT_FOUND)

    invoice_number = online_invoice_number(reservation)
    response = HttpResponse(build_online_invoice_pdf(reservation, invoice_number), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="Facture_WilloBarber_{invoice_number}.pdf"'
    return response


@api_view(['GET'])
@permission_classes([IsStaffRole])
def download_final_invoice(request, pk):
    try:
        reservation = Reservation.objects.get(pk=pk)
    except Reservation.DoesNotExist:
        return Response({'error': 'Réservation introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    if reservation.payment_status != 'paid_onsite':
        return Response(
            {'error': "La réservation n'est pas marquée comme payée au salon."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    pdf_bytes, invoice_number = build_final_invoice_pdf(reservation)
    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{invoice_number}.pdf"'
    return response


@api_view(['POST'])
@permission_classes([IsStaffRole])
def upload_client_media(request):
    uploaded_file = request.FILES.get('file')
    if not uploaded_file:
        return Response({'error': 'Fichier requis.'}, status=status.HTTP_400_BAD_REQUEST)

    client_id = request.data.get('client')
    if not client_id:
        return Response({'error': 'client requis.'}, status=status.HTTP_400_BAD_REQUEST)

    try:
        client = User.objects.get(pk=client_id)
    except User.DoesNotExist:
        return Response({'error': 'Client introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    media_type = request.data.get('media_type') or (
        'video' if (uploaded_file.content_type or '').startswith('video') else 'photo'
    )
    if media_type not in ('photo', 'video'):
        return Response({'error': 'media_type doit être "photo" ou "video".'}, status=status.HTTP_400_BAD_REQUEST)

    reservation = None
    reservation_id = request.data.get('reservation')
    if reservation_id:
        try:
            reservation = Reservation.objects.get(pk=reservation_id)
        except Reservation.DoesNotExist:
            return Response({'error': 'Réservation introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    try:
        result = cloudinary.uploader.upload(
            uploaded_file,
            resource_type='video' if media_type == 'video' else 'image',
            folder='willobarber/client_media',
        )
        logger.info(f"[CLOUDINARY UPLOAD SUCCESS] public_id={result.get('public_id')} url={result.get('secure_url')}")
    except Exception as e:
        logger.error(f"[CLOUDINARY UPLOAD ERROR] {e}")
        return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    media = ClientMedia.objects.create(
        client=client,
        reservation=reservation,
        media_type=media_type,
        cloudinary_url=result.get('secure_url'),
        cloudinary_public_id=result.get('public_id'),
        caption=request.data.get('caption', ''),
    )

    return Response(ClientMediaSerializer(media).data, status=status.HTTP_201_CREATED)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def upload_profile_picture(request):
    uploaded_file = request.FILES.get('file')
    if not uploaded_file:
        return Response({'error': 'Fichier requis.'}, status=status.HTTP_400_BAD_REQUEST)

    try:
        result = cloudinary.uploader.upload(
            uploaded_file,
            resource_type='image',
            folder='willobarber/profiles',
        )
        logger.info(f"[CLOUDINARY UPLOAD SUCCESS] public_id={result.get('public_id')} url={result.get('secure_url')}")
    except Exception as e:
        logger.error(f"[CLOUDINARY UPLOAD ERROR] {e}")
        return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    request.user.profile_picture = result.get('secure_url')
    request.user.save(update_fields=['profile_picture'])

    return Response(UserSerializer(request.user).data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def client_media_list(request, pk):
    is_staff = request.user.is_superuser or getattr(request.user, 'role', None) == 'admin'
    if not is_staff and request.user.id != pk:
        return Response({'error': "Vous n'avez pas accès à cette galerie."}, status=status.HTTP_403_FORBIDDEN)

    media = ClientMedia.objects.filter(client_id=pk)
    if not is_staff:
        media = media.filter(shared_with_client=True)
    return Response(ClientMediaSerializer(media, many=True).data)


@api_view(['PATCH'])
@permission_classes([IsStaffRole])
def share_client_media(request, pk):
    try:
        media = ClientMedia.objects.get(pk=pk)
    except ClientMedia.DoesNotExist:
        return Response({'error': 'Media introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    media.shared_with_client = True
    media.save(update_fields=['shared_with_client'])
    return Response(ClientMediaSerializer(media).data)


def _serialize_clients(queryset):
    clients = queryset.annotate(
        total_reservations=Count('reservation', distinct=True),
        last_visit=Max('reservation__date'),
        total_spent=Sum('reservation__amount_paid'),
    ).order_by('-date_joined')

    return [
        {
            'id': c.id,
            'first_name': c.first_name,
            'last_name': c.last_name,
            'email': c.email,
            'phone': c.phone,
            'total_reservations': c.total_reservations,
            'last_visit': c.last_visit.isoformat() if c.last_visit else None,
            'total_spent': float(c.total_spent or 0),
            'date_joined': c.date_joined.isoformat(),
            'status': c.status,
            'is_archived': c.is_archived,
            'deleted_at': c.deleted_at.isoformat() if c.deleted_at else None,
            'deletion_reason': c.deletion_reason,
        }
        for c in clients
    ]


@api_view(['GET'])
@permission_classes([IsStaffRole])
def client_list(request):
    clients = User.objects.filter(role='client', is_deleted=False, is_archived=False)
    return Response(_serialize_clients(clients))


@api_view(['GET'])
@permission_classes([IsStaffRole])
def deleted_client_list(request):
    # L'onglet "Archivés" regroupe les clients archivés par le coiffeur et les
    # comptes supprimés (staff ou RGPD) : tous se reprennent via restore_client.
    clients = User.objects.filter(role='client').filter(Q(is_deleted=True) | Q(is_archived=True))
    return Response(_serialize_clients(clients))


@api_view(['DELETE'])
@permission_classes([IsStaffRole])
def delete_client(request, pk):
    try:
        client = User.objects.get(pk=pk, role='client', is_deleted=False)
    except User.DoesNotExist:
        return Response({'error': 'Client introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    client.is_deleted = True
    client.deleted_at = timezone.now()
    client.save(update_fields=['is_deleted', 'deleted_at'])
    return Response(status=status.HTTP_204_NO_CONTENT)


@api_view(['POST'])
@permission_classes([IsStaffRole])
def restore_client(request, pk):
    try:
        client = User.objects.get(Q(is_deleted=True) | Q(is_archived=True), pk=pk, role='client')
    except User.DoesNotExist:
        return Response({'error': 'Client supprimé introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    client.is_archived = False
    client.is_deleted = False
    client.deleted_at = None
    client.deletion_reason = ''
    # Réactive la connexion pour une auto-suppression restaurée (delete_client
    # côté staff ne touche jamais is_active, donc ce champ ne peut valoir False
    # ici que si le compte a été anonymisé via /auth/delete-account/).
    client.is_active = True
    client.save(update_fields=['is_archived', 'is_deleted', 'deleted_at', 'deletion_reason', 'is_active'])
    return Response(_serialize_clients(User.objects.filter(pk=client.pk))[0])


def _set_client_archived(pk, archived):
    try:
        client = User.objects.get(pk=pk, role='client', is_deleted=False)
    except User.DoesNotExist:
        return Response({'error': 'Client introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    client.is_archived = archived
    client.save(update_fields=['is_archived'])
    return Response(_serialize_clients(User.objects.filter(pk=client.pk))[0])


@api_view(['PATCH'])
@permission_classes([IsStaffRole])
def archive_client(request, pk):
    return _set_client_archived(pk, True)


@api_view(['PATCH'])
@permission_classes([IsStaffRole])
def unarchive_client(request, pk):
    return _set_client_archived(pk, False)


@api_view(['PATCH'])
@permission_classes([IsStaffRole])
def update_client_status(request, pk):
    new_status = request.data.get('status')
    if new_status not in dict(User.STATUS_CHOICES):
        return Response({'error': 'Statut invalide.'}, status=status.HTTP_400_BAD_REQUEST)

    try:
        client = User.objects.get(pk=pk, role='client', is_deleted=False)
    except User.DoesNotExist:
        return Response({'error': 'Client introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    client.status = new_status
    client.save(update_fields=['status'])
    return Response(_serialize_clients(User.objects.filter(pk=client.pk))[0])


@api_view(['GET', 'PATCH'])
@permission_classes([IsStaffRole])
def notification_preferences(request):
    prefs, _ = NotificationPreference.objects.get_or_create(user=request.user)

    if request.method == 'PATCH':
        serializer = NotificationPreferenceSerializer(prefs, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data)

    return Response(NotificationPreferenceSerializer(prefs).data)


@api_view(['GET', 'POST'])
@permission_classes([IsStaffRole])
def payment_cards(request):
    if request.method == 'POST':
        is_first_card = not request.user.payment_cards.exists()
        serializer = PaymentCardSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        serializer.save(owner=request.user, is_primary=is_first_card)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    cards = request.user.payment_cards.all()
    return Response(PaymentCardSerializer(cards, many=True).data)


@api_view(['PATCH'])
@permission_classes([IsStaffRole])
def payment_card_set_primary(request, pk):
    try:
        card = PaymentCard.objects.get(pk=pk, owner=request.user)
    except PaymentCard.DoesNotExist:
        return Response({'error': 'Carte introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    with transaction.atomic():
        request.user.payment_cards.exclude(pk=card.pk).update(is_primary=False)
        card.is_primary = True
        card.save(update_fields=['is_primary'])

    return Response(PaymentCardSerializer(card).data)


@api_view(['DELETE'])
@permission_classes([IsStaffRole])
def payment_card_delete(request, pk):
    try:
        card = PaymentCard.objects.get(pk=pk, owner=request.user)
    except PaymentCard.DoesNotExist:
        return Response({'error': 'Carte introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    card.delete()
    return Response(status=status.HTTP_204_NO_CONTENT)


@api_view(['GET', 'PATCH'])
@permission_classes([IsStaffRole])
def commission_settings(request):
    shop = Barbershop.objects.filter(owner=request.user).first() or Barbershop.objects.first()
    if not shop:
        return Response({'error': 'Aucun salon configuré.'}, status=status.HTTP_404_NOT_FOUND)

    if request.method == 'PATCH':
        rate = request.data.get('commission_rate')
        try:
            shop.commission_rate = round(float(rate), 1)
        except (TypeError, ValueError):
            return Response({'error': 'Taux de commission invalide.'}, status=status.HTTP_400_BAD_REQUEST)
        shop.save(update_fields=['commission_rate'])

    return Response({'commission_rate': float(shop.commission_rate)})


@api_view(['GET'])
@permission_classes([IsStaffRole])
def security_settings(request):
    return Response({'two_factor_enabled': request.user.two_factor_enabled})


@api_view(['PATCH'])
@permission_classes([IsStaffRole])
def toggle_two_factor(request):
    enabled = request.data.get('enabled')
    if not isinstance(enabled, bool):
        return Response({'error': "Le champ 'enabled' doit être un booléen."}, status=status.HTTP_400_BAD_REQUEST)
    request.user.two_factor_enabled = enabled
    request.user.save(update_fields=['two_factor_enabled'])
    return Response({'two_factor_enabled': request.user.two_factor_enabled})


@api_view(['GET'])
@permission_classes([IsStaffRole])
def session_list(request):
    auth_header = request.META.get('HTTP_AUTHORIZATION', '')
    current_access_token = auth_header.split(' ')[-1] if auth_header else ''

    sessions = request.user.sessions.all()
    return Response([
        {
            'id': s.id,
            'device': s.device or 'Appareil inconnu',
            'created_at': s.created_at.isoformat(),
            'current': s.access_token == current_access_token,
        }
        for s in sessions
    ])


@api_view(['DELETE'])
@permission_classes([IsStaffRole])
def session_delete(request, pk):
    try:
        session = UserSession.objects.get(pk=pk, user=request.user)
    except UserSession.DoesNotExist:
        return Response({'error': 'Session introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    try:
        RefreshToken(session.refresh_token).blacklist()
    except Exception:
        pass
    session.delete()
    return Response(status=status.HTTP_204_NO_CONTENT)


@api_view(['POST'])
@permission_classes([IsStaffRole])
def send_reminders(request):
    result = send_reminders_for_tomorrow()
    return Response(result)


class ReviewListCreateView(APIView):
    """GET (Willo) liste tous les avis · POST (client) soumet un avis sur son RDV terminé."""

    def get_permissions(self):
        if self.request.method == 'POST':
            return [IsAuthenticated()]
        return [IsStaffRole()]

    def get(self, request):
        reviews = Review.objects.select_related('user', 'reservation__service').all()
        return Response(ReviewSerializer(reviews, many=True).data)

    def post(self, request):
        reservation_id = request.data.get('reservation')
        try:
            reservation = Reservation.objects.get(pk=reservation_id)
        except (Reservation.DoesNotExist, TypeError, ValueError):
            return Response({'error': 'Réservation introuvable.'}, status=status.HTTP_404_NOT_FOUND)

        if reservation.user_id != request.user.id:
            return Response({'error': "Cette réservation ne vous appartient pas."}, status=status.HTTP_403_FORBIDDEN)

        if reservation.status != 'completed':
            return Response(
                {'error': 'Seule une réservation terminée peut être notée.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if Review.objects.filter(reservation=reservation).exists():
            return Response(
                {'error': 'Un avis a déjà été soumis pour cette réservation.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            rating = int(request.data.get('rating'))
        except (TypeError, ValueError):
            return Response({'error': 'Note invalide.'}, status=status.HTTP_400_BAD_REQUEST)
        if rating < 1 or rating > 5:
            return Response({'error': 'La note doit être comprise entre 1 et 5.'}, status=status.HTTP_400_BAD_REQUEST)

        review = Review.objects.create(
            user=request.user,
            reservation=reservation,
            rating=rating,
            comment=request.data.get('comment') or '',
        )
        return Response(ReviewSerializer(review).data, status=status.HTTP_201_CREATED)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def my_reviews(request):
    reviews = Review.objects.filter(user=request.user).select_related('reservation__service')
    return Response(ReviewSerializer(reviews, many=True).data)


@api_view(['POST'])
@permission_classes([IsStaffRole])
def reply_to_review(request, pk):
    try:
        review = Review.objects.get(pk=pk)
    except Review.DoesNotExist:
        return Response({'error': 'Avis introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    reply_text = (request.data.get('reply') or '').strip()
    if not reply_text:
        return Response({'error': 'La réponse ne peut pas être vide.'}, status=status.HTTP_400_BAD_REQUEST)

    review.reply = reply_text
    review.replied_at = timezone.now()
    review.save(update_fields=['reply', 'replied_at'])

    return Response(ReviewSerializer(review).data)


def landing_page(request):
    return TemplateResponse(request, 'landing.html', {})


def qr_code(request):
    url = "https://willobarber-production-6951.up.railway.app"
    qr = qrcode.QRCode(version=1, box_size=10, border=4)
    qr.add_data(url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    buffer.seek(0)
    return HttpResponse(buffer, content_type="image/png")


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def redeem_loyalty(request):
    """Déduit 1000 pts du client et confirme le rachat d'une coupe gratuite."""
    user = request.user
    REDEEM_COST = 1000
    if (user.loyalty_points or 0) < REDEEM_COST:
        return Response(
            {'error': f'Solde insuffisant. {REDEEM_COST - (user.loyalty_points or 0)} pts manquants.'},
            status=400
        )
    user.loyalty_points -= REDEEM_COST
    user.save(update_fields=['loyalty_points'])
    return Response({'loyalty_points': user.loyalty_points, 'redeemed': REDEEM_COST})
